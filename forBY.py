
import os
from docx import *
import asyncio
from pyppeteer import launch
import xlwt
from urllib.request import urlretrieve
import os
import shutil
import hashlib
import androguard.core.bytecodes.apk as apk
from androguard.util import get_certificate_name_string
from asn1crypto import x509
import time

##### 额外的py包 pyppeteer、xlwt、androguard、python-docx

__cols = ["研判时间", "银行类别", "程序名称", "样本MD5", "应用市场", "下载链接", "是否可以下载", "是否为仿冒",
                "恶意行为描述", "备注", "是否存活", "详情页", "签名", "签名MD5"]

###### 指定某银行
XXbank = "光大银行"

###### 指定扫描的docx文件
XXdocx = "副本.docx"

workbook = xlwt.Workbook(encoding="utf8")
worksheet = workbook.add_sheet("验证结果", cell_overwrite_ok=True)
for i,_ in enumerate(__cols):
    worksheet.write(0, i, __cols[i])


def get_all_info(path):
    all_info = []
    doc = Document(path)
    print("docx中的应用商店、发布页面、下载链接信息如下\n")
    for i, table in enumerate(doc.tables):
        for j, row in enumerate(table.rows):
            hasfind = False
            for k, cell in enumerate(row.cells):
                cellText = cell.text.strip()
                if cellText.find("http") == 0:   ###找到链接了, 标志位置True
                    hasfind = True
            if hasfind is True:
                shopname = row.cells[1].text.splitlines()[0].strip().split("(")[0]
                shopurl = row.cells[3].text.strip()
                downloadurl = row.cells[4].text.strip()
                print(shopname, shopurl, downloadurl)
                all_info.append([shopname, shopurl, downloadurl])
    print("docx有效信息收集完毕\n")
    return all_info


async def main():

    all_info = get_all_info("副本.docx")
    await asyncio.sleep(2)
    browser = await launch(headless=False)  # headless=False,devtools=True
    shoppage = await  browser.newPage()
    await shoppage.setViewport({'width': 1000, 'height': 1200})

##### 登记发布页面是否存在
    print("选择题开始，请谨慎选择！！\n")
    for i, info in enumerate(all_info):
        worksheet.write(i + 1, 0, time.strftime('%Y-%m-%d'))
        worksheet.write(i + 1, 1, XXbank)
        try:
            await shoppage.goto(info[1], waitUntil='documentloaded')
        except:
            worksheet.write(i+1, 11, "无")
            continue
        await asyncio.sleep(0.5)
        await shoppage.addScriptTag(path="inject.js")
        isAvailable = await shoppage.evaluate("() => window.__shi")
        isAvailable = info[1] if isAvailable else "无"
        print(isAvailable)
        worksheet.write(i + 1, 11, isAvailable)
    await shoppage.close()
    await browser.close()
    print("选择题结束，开始逐个下载\n")


    if os.path.exists("out"):
        shutil.rmtree("out", ignore_errors=True)
    else:
        os.mkdir("out")
    await asyncio.sleep(1)

    for i, info in enumerate(all_info):
        print("开始下载:"+ info[2])
        canDown = True
        try:
            urlretrieve(info[2], "./out/cur.apk")
        except:
            canDown = False
            print("！！！该链接无法下载！！！\n")
            worksheet.write(i + 1, 5, info[2])
            worksheet.write(i + 1, 6, "！！！待定")
            worksheet.write(i + 1, 9, "下载链接待人工确定")
        if canDown:
            with open("./out/cur.apk", "rb") as f:
                file = f.read()
                curapk = apk.APK(file, True)
                appname = curapk.get_app_name()
                certs = set(curapk.get_certificates_der_v2() + [curapk.get_certificate_der(x) for x in
                                                                curapk.get_signature_names()])
                Issuer, signMd5 = "", ""
                for cert in certs:
                    x509_cert = x509.Certificate.load(cert)
                    Issuer = get_certificate_name_string(x509_cert.issuer.native, short=True)
                    signMd5 = hashlib.md5(cert).hexdigest().upper()
                apkmd5 = hashlib.md5(file).hexdigest().upper()
                ###文件命名方式： 第几行_应用商店_MD5.apk，方便检索、对比
                newname= "./out/" + str(i+2) + "_" + info[0] + "_" + apkmd5 + ".apk"
                print("######  当前apk信息  #######")
                print("\t应用名:" + appname)
                print("\tMD5:" + apkmd5)
                print("\t签名:" + Issuer)
                print("\t签名MD5:" + signMd5)
                print()
                os.rename("./out/cur.apk", newname)
                await asyncio.sleep(0.3)
                worksheet.write(i + 1, 2, appname)
                worksheet.write(i + 1, 3, apkmd5)
                worksheet.write(i + 1, 4, info[0])
                worksheet.write(i + 1, 5, info[2])
                worksheet.write(i + 1, 6, "是")
                worksheet.write(i + 1, 10, "是")
                worksheet.write(i + 1, 12, Issuer)
                worksheet.write(i + 1, 13, signMd5)

    workbook.save("最终结果.xls")


asyncio.get_event_loop().run_until_complete(main())
